# utils/data_reader.py

import openpyxl


def read_valid_credentials(filepath: str = "validdata.xlsx", row: int = 2):
    """
    Reads username and password from an Excel file.

    Layout expected in the workbook:
        Row 1  → Header  (TC_ID | Username | Password)
        Row 2+ → Data rows

    Default row=2  →  standard_user / secret_sauce

    Args:
        filepath (str): Path to the Excel test-data file.
        row      (int): Row number to read (1-indexed; row 1 = header).

    Returns:
        tuple[str, str]: (username, password)
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    username = ws.cell(row=row, column=2).value
    password = ws.cell(row=row, column=3).value
    return username, password
